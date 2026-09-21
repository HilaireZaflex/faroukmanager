"""
Routes API — « Missions d'appels ».
Prefix: /api/missions-appels

Créées par l'encadrement (Admin, RC, Manager, Conformité, Responsable Produit
& Qualité), exécutées par les téléconseillères.
"""
from fastapi import APIRouter, Depends, Body, Query, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import func
from io import BytesIO
from typing import Optional

from app.core.database import get_db
from app.api.routes.auth import get_current_user
from app.models.user import User
from app.models.pdv import PDV
from app.models.mission_appel import MissionAppel, MissionCible, StatutMission
from app.models.appel_tc import STATUT_LABELS
from app.services import mission_appel_service as svc

router = APIRouter()


# ─────────────────────────────────────────────────────────────────────────────
# DONNÉES DE RÉFÉRENCE (formulaire de création)
# ─────────────────────────────────────────────────────────────────────────────

@router.get("/missions-appels/filtres")
def valeurs_filtres(db: Session = Depends(get_db), _: User = Depends(get_current_user)):
    """Valeurs disponibles pour composer une mission."""
    sup = [r[0] for r in db.query(PDV.superviseur).filter(
        PDV.superviseur.isnot(None), PDV.superviseur != '').distinct().order_by(PDV.superviseur).all()]
    gest = [r[0] for r in db.query(PDV.gestionnaire).filter(
        PDV.gestionnaire.isnot(None), PDV.gestionnaire != '').distinct().order_by(PDV.gestionnaire).all()]
    zones = [r[0] for r in db.query(PDV.zone).filter(
        PDV.zone.isnot(None), PDV.zone != '').distinct().order_by(PDV.zone).all()]
    quartiers = [r[0] for r in db.query(PDV.quartier).filter(
        PDV.quartier.isnot(None), PDV.quartier != '').distinct().order_by(PDV.quartier).all()]

    roles_personnes = [r[0] for r in db.query(func.lower(User.role)).filter(
        User.is_active == True).distinct().order_by(func.lower(User.role)).all()]  # noqa: E712

    # Période de référence par défaut : le dernier mois disposant de données,
    # sinon le mois courant (évite de composer une mission sur un mois vide).
    from datetime import date as _date
    from app.models.performance import MonthlyPerformance
    auj = _date.today()
    dernier = db.query(MonthlyPerformance.annee, MonthlyPerformance.mois).order_by(
        MonthlyPerformance.annee.desc(), MonthlyPerformance.mois.desc()).first()
    annee_defaut, mois_defaut = (dernier[0], dernier[1]) if dernier else (auj.year, auj.month)

    return {
        "annee_defaut": annee_defaut,
        "mois_defaut": mois_defaut,
        "superviseurs": sup,
        "gestionnaires": gest,
        "zones": zones,
        "quartiers": quartiers,
        "situations": [
            {"code": "OMY_INACTIF",    "label": "OMY inactif"},
            {"code": "OMY_BAISSE",     "label": "OMY en baisse (≥ 30 %)"},
            {"code": "NAFAMA_INACTIF", "label": "NAFAMA inactif"},
            {"code": "NAFAMA_BAISSE",  "label": "NAFAMA en baisse (≥ 30 %)"},
            {"code": "KAABU_INACTIF",  "label": "KAABU inactif"},
            {"code": "JAMAIS_APPELE",  "label": "Jamais appelé"},
            {"code": "PAS_APPELE_DEPUIS", "label": "Pas appelé depuis X jours"},
        ],
        "roles_personnes": roles_personnes,
        "teleconseilleres": svc.tcs_disponibles(db),
        "type_missions": [
            {"code": "RELANCE_ACTIVITE",     "label": "📈 Relance activité"},
            {"code": "VERIFICATION_TERRAIN", "label": "🚶 Vérification terrain"},
            {"code": "ENQUETE",              "label": "📊 Enquête / sondage"},
            {"code": "MIGRATION",            "label": "🚀 Migration"},
            {"code": "CONTROLE_CONFORMITE",  "label": "🛡️ Contrôle conformité"},
            {"code": "AUTRE",                "label": "📞 Autre"},
        ],
        "priorites": [
            {"code": "NORMALE", "label": "Normale"},
            {"code": "HAUTE",   "label": "🔶 Haute"},
            {"code": "URGENTE", "label": "🔴 Urgente"},
        ],
        "statuts_appel": [
            {"code": k, "label": v} for k, v in STATUT_LABELS.items()
        ],
    }


@router.get("/missions-appels/personnes")
def personnes_cibles(role: Optional[str] = None, q: Optional[str] = None,
                     db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    """Personnes pouvant être appelées (superviseurs, gestionnaires, autres)."""
    svc.exige_createur(current_user)
    query = db.query(User).filter(User.is_active == True)  # noqa: E712
    if role:
        query = query.filter(func.lower(User.role) == role.lower())
    users = query.order_by(User.nom).limit(1000).all()
    out = []
    for u in users:
        nom = svc._nom_complet(u)
        if q and q.lower() not in nom.lower():
            continue
        out.append({
            "user_id": u.id, "nom": nom, "role": svc.normalise_role(u),
            "telephone": u.telephone, "zone": u.zone,
        })
    return out


@router.get("/missions-appels/pdv-candidats")
def pdv_candidats(annee: Optional[int] = None, mois: Optional[int] = None,
                  db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    """Tous les PDV actifs (liste simple affichée dans l'assistant)."""
    svc.exige_createur(current_user)
    return svc.liste_pdv_candidats(db, annee, mois)


@router.post("/missions-appels/apercu")
def apercu(payload: dict = Body(...), db: Session = Depends(get_db),
           current_user: User = Depends(get_current_user)):
    svc.exige_createur(current_user)
    return svc.apercu_cibles(
        db, payload.get("filtres") or {},
        payload.get("annee"), payload.get("mois"),
    )


# ─────────────────────────────────────────────────────────────────────────────
# CÔTÉ TÉLÉCONSEILLÈRE (déclaré AVANT /{mission_id} pour éviter les conflits)
# ─────────────────────────────────────────────────────────────────────────────

@router.get("/missions-appels/mes-missions")
def mes_missions(db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    return svc.mes_missions(db, current_user)


@router.get("/missions-appels/ma-file")
def ma_file(db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    return svc.cibles_a_faire(db, current_user)


@router.post("/missions-appels/cibles/{cible_id}/appel")
def enregistrer_appel(cible_id: int, payload: dict = Body(...),
                      db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    return svc.enregistrer_appel(db, cible_id, payload, current_user)


@router.patch("/missions-appels/cibles/{cible_id}")
def marquer_cible(cible_id: int, payload: dict = Body(...),
                  db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    return svc.marquer_cible(db, cible_id, payload, current_user)


# ─────────────────────────────────────────────────────────────────────────────
# CRÉATION / PILOTAGE (encadrement)
# ─────────────────────────────────────────────────────────────────────────────

@router.post("/missions-appels")
def creer(payload: dict = Body(...), db: Session = Depends(get_db),
          current_user: User = Depends(get_current_user)):
    m = svc.creer_mission(db, payload, current_user)
    return {"id": m.id, "titre": m.titre, "statut": m.statut.value}


@router.get("/missions-appels")
def lister(statut: Optional[str] = None, actives: bool = False, mes_missions: bool = False,
           db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    return svc.lister_missions(db, current_user, statut=statut,
                               seulement_actives=actives, mes_missions=mes_missions)


@router.get("/missions-appels/{mission_id}")
def detail(mission_id: int, db: Session = Depends(get_db),
           current_user: User = Depends(get_current_user)):
    # Une TC ne voit que ses propres cibles ; l'encadrement voit tout
    seulement_miennes = not svc.peut_creer(current_user)
    return svc.detail_mission(db, mission_id, current_user, seulement_mes_cibles=seulement_miennes)


@router.patch("/missions-appels/{mission_id}")
def modifier(mission_id: int, payload: dict = Body(...), db: Session = Depends(get_db),
             current_user: User = Depends(get_current_user)):
    m = svc.modifier_mission(db, mission_id, payload, current_user)
    return {"id": m.id, "statut": m.statut.value}


@router.post("/missions-appels/{mission_id}/attribuer")
def attribuer(mission_id: int, payload: dict = Body(...), db: Session = Depends(get_db),
              current_user: User = Depends(get_current_user)):
    return svc.attribuer(db, mission_id, payload, current_user)


@router.post("/missions-appels/{mission_id}/desattribuer")
def desattribuer(mission_id: int, payload: dict = Body(...), db: Session = Depends(get_db),
                 current_user: User = Depends(get_current_user)):
    return svc.desattribuer(db, mission_id, payload, current_user)


@router.post("/missions-appels/{mission_id}/cloturer")
def cloturer(mission_id: int, payload: dict = Body(default={}), db: Session = Depends(get_db),
             current_user: User = Depends(get_current_user)):
    statut = (payload or {}).get("statut") or "TERMINEE"
    m = svc.cloturer_mission(db, mission_id, current_user, statut)
    return {"id": m.id, "statut": m.statut.value}


@router.get("/missions-appels/{mission_id}/stats")
def stats(mission_id: int, db: Session = Depends(get_db),
          current_user: User = Depends(get_current_user)):
    svc.exige_createur(current_user)
    return svc.stats_mission(db, mission_id, current_user)


@router.get("/missions-appels/{mission_id}/export")
def export(mission_id: int, db: Session = Depends(get_db),
           current_user: User = Depends(get_current_user)):
    """Export Excel du suivi d'une mission."""
    svc.exige_createur(current_user)
    data = svc.detail_mission(db, mission_id, current_user)
    m = data["mission"]

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    ws = wb.active
    ws.title = f"Mission {mission_id}"

    entetes = ["Type", "Cible", "N° PDV", "Quartier", "Superviseur", "Téléphone",
               "Motif", "Téléconseillère", "Statut", "Dernier statut d'appel",
               "Nb appels", "Dernier appel", "Motif d'abandon"]
    fill = PatternFill("solid", fgColor="FF6900")
    font = Font(bold=True, color="FFFFFF")
    for i, h in enumerate(entetes, 1):
        c = ws.cell(row=1, column=i, value=h)
        c.fill = fill
        c.font = font
        c.alignment = Alignment(horizontal="center")

    for c in data["cibles"]:
        ws.append([
            c.get("type_cible"), c.get("nom"), c.get("pdv_numero") or "",
            c.get("quartier") or "", c.get("superviseur") or "", c.get("telephone") or "",
            c.get("motif") or "", c.get("assigned_to_nom") or "", c.get("statut") or "",
            c.get("dernier_statut_label") or "", c.get("nb_appels") or 0,
            (c.get("dernier_appel_at") or "")[:19].replace("T", " "),
            c.get("abandon_motif") or "",
        ])

    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = max(
            (len(str(c.value)) for c in col if c.value), default=10) + 2

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    nom = "".join(ch for ch in (m.get("titre") or "mission") if ch.isalnum() or ch in " -_")[:40]
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="mission_{nom}.xlsx"'},
    )
