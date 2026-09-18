"""
Routes API Appels Migration — prospection « migration en commission directe ».

Les téléconseillères appellent leurs PDV de type RS et KIOSQUE et posent trois
questions. Un PDV est VALIDÉ seulement si les trois critères sont remplis.
L'encadrement (admin, RC, superviseurs) consulte la liste et peut l'exporter.
"""
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import func, or_
from typing import Optional, List
from datetime import datetime
from pydantic import BaseModel
import io

from app.core.database import get_db
from app.api.routes.auth import get_current_user
from app.api.routes.appels_tc import _full_name, _pdv_autorise_pour_tc
from app.models.user import User
from app.models.pdv import PDV
from app.models.appel_migration import (
    AppelMigration, MIGRATION_TYPES_PDV, MIGRATION_TYPES_PIECE,
    MIGRATION_TYPE_PIECE_LABELS,
)

router = APIRouter()


# ─── Helpers ─────────────────────────────────────────────────────────────────

def _role(user: User) -> str:
    return str(user.role or '').lower().replace('userrole.', '').strip()


def _est_encadrement(user: User) -> bool:
    """Qui peut consulter / exporter la liste des appels migration."""
    return _role(user) in (
        'admin', 'manager', 'rc', 'superviseur', 'conformite',
        'responsable_produit_et_qualit_oprationnelle_',
    )


def _est_admin(user: User) -> bool:
    return _role(user) in ('admin', 'manager')


class MigrationIn(BaseModel):
    numero_pdv: str
    nom_pdv: Optional[str] = None
    type_pdv: Optional[str] = None
    veut_migrer: bool = False
    a_rccm: bool = False
    a_piece_identite: bool = False
    type_piece: Optional[str] = None
    commentaire: Optional[str] = None


def _fmt(m: AppelMigration) -> dict:
    return {
        "id": m.id,
        "numero_pdv": m.numero_pdv,
        "nom_pdv": m.nom_pdv,
        "type_pdv": m.type_pdv,
        "tc_user_id": m.tc_user_id,
        "tc_nom": m.tc_nom,
        "veut_migrer": bool(m.veut_migrer),
        "a_rccm": bool(m.a_rccm),
        "a_piece_identite": bool(m.a_piece_identite),
        "type_piece": m.type_piece,
        "type_piece_label": MIGRATION_TYPE_PIECE_LABELS.get(m.type_piece or "", m.type_piece) if m.type_piece else None,
        "statut": m.statut,
        "motif_rejet": m.motif_rejet,
        "commentaire": m.commentaire,
        "pieces_au_bureau": bool(getattr(m, "pieces_au_bureau", False)),
        "date_depot_bureau": m.date_depot_bureau.isoformat() if getattr(m, "date_depot_bureau", None) else None,
        "depot_par_nom": getattr(m, "depot_par_nom", None),
        "created_at": m.created_at.isoformat() if m.created_at else None,
    }


def _evaluer(veut_migrer: bool, a_rccm: bool, a_piece: bool):
    """Retourne (statut, motif_rejet)."""
    manques = []
    if not veut_migrer:
        manques.append("ne souhaite pas migrer en commission directe")
    if not a_rccm:
        manques.append("RCCM manquant")
    if not a_piece:
        manques.append("pièce d'identité valide manquante")
    if manques:
        return "REJETE", "Critère(s) non rempli(s) : " + ", ".join(manques)
    return "VALIDE", None


def _derniers_appels(db: Session, numeros: List[str]) -> dict:
    """Dernier appel migration enregistré pour chaque PDV fourni."""
    derniers = {}
    if not numeros:
        return derniers
    for m in db.query(AppelMigration).filter(
        AppelMigration.numero_pdv.in_(numeros)
    ).order_by(AppelMigration.created_at.asc()).all():
        derniers[m.numero_pdv] = m
    return derniers


# ─── File d'appels migration de la téléconseillère ───────────────────────────

@router.get("/tc/migration/pdv")
def list_pdv_migration(
    tc_user_id: Optional[int] = Query(None, description="Filtrer sur un compte TC (encadrement)"),
    type_pdv: Optional[str] = Query(None, description="RS ou KIOSQUE"),
    search: Optional[str] = Query(None),
    deja_appeles: Optional[bool] = Query(None, description="True = déjà appelés, False = restants"),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """PDV de type RS / KIOSQUE à appeler pour la migration.

    Une téléconseillère ne voit que SES PDV ; l'encadrement voit tout (filtrable).
    """
    q = db.query(PDV).filter(
        PDV.statut == 'ACTIF',
        PDV.type_pdv.in_(MIGRATION_TYPES_PDV),
    )

    role = _role(current_user)
    if role in ('teleconseillere', 'tc'):
        q = q.filter(or_(
            PDV.teleconseillere_user_id == current_user.id,
            PDV.teleconseillere == _full_name(current_user),
        ))
    elif tc_user_id:
        q = q.filter(PDV.teleconseillere_user_id == tc_user_id)

    if type_pdv and type_pdv in MIGRATION_TYPES_PDV:
        q = q.filter(PDV.type_pdv == type_pdv)

    if search:
        like = f"%{search}%"
        q = q.filter(or_(
            PDV.numero_pdv.ilike(like),
            PDV.nom.ilike(like),
            PDV.quartier.ilike(like),
        ))

    pdvs = q.all()
    derniers = _derniers_appels(db, [p.numero_pdv for p in pdvs])

    items = []
    for p in pdvs:
        m = derniers.get(p.numero_pdv)
        items.append({
            "numero_pdv": p.numero_pdv,
            "nom": p.nom,
            "type_pdv": p.type_pdv,
            "zone": p.zone,
            "sous_zone": p.sous_zone,
            "quartier": p.quartier,
            "telephone": p.telephone,
            "numero_personnel": p.numero_personnel,
            "nom_gerant": p.nom_gerant,
            "teleconseillere": p.teleconseillere,
            "deja_appele": m is not None,
            "dernier_statut": m.statut if m else None,
            "dernier_appel": m.created_at.isoformat() if m and m.created_at else None,
            # Infos du dernier appel, utiles pour cocher le dépôt des pièces
            "dernier_appel_id": m.id if m else None,
            "veut_migrer": bool(m.veut_migrer) if m else None,
            "a_rccm": bool(m.a_rccm) if m else None,
            "a_piece_identite": bool(m.a_piece_identite) if m else None,
            "type_piece": m.type_piece if m else None,
            "pieces_au_bureau": bool(getattr(m, "pieces_au_bureau", False)) if m else False,
            "date_depot_bureau": (
                m.date_depot_bureau.isoformat()
                if m and getattr(m, "date_depot_bureau", None) else None
            ),
        })

    if deja_appeles is True:
        items = [i for i in items if i["deja_appele"]]
    elif deja_appeles is False:
        items = [i for i in items if not i["deja_appele"]]

    items.sort(key=lambda x: (x["deja_appele"], x["type_pdv"] or "", x["numero_pdv"] or ""))

    return {
        "total": len(items),
        "total_pdv": len(pdvs),
        "deja_appeles": sum(1 for i in items if i["deja_appele"]),
        "items": items,
        "types": MIGRATION_TYPES_PDV,
    }


@router.get("/tc/migration/pdv/{numero_pdv}")
def historique_pdv(
    numero_pdv: str,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Historique des appels migration d'un PDV."""
    pdv = db.query(PDV).filter(PDV.numero_pdv == numero_pdv).first()
    if pdv is not None and not _pdv_autorise_pour_tc(db, current_user, numero_pdv):
        raise HTTPException(status_code=403, detail="Ce point de vente est affecté à une autre téléconseillère")

    appels = db.query(AppelMigration).filter(
        AppelMigration.numero_pdv == numero_pdv
    ).order_by(AppelMigration.created_at.desc()).limit(50).all()
    return {"items": [_fmt(a) for a in appels]}


# ─── Enregistrement d'un appel migration ─────────────────────────────────────

@router.post("/tc/migration")
def create_appel_migration(
    body: MigrationIn,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Enregistrer un appel migration (les 3 critères déterminent VALIDE / REJETE)."""
    numero = (body.numero_pdv or "").strip()
    if not numero:
        raise HTTPException(status_code=400, detail="Numéro de PDV obligatoire")

    pdv = db.query(PDV).filter(PDV.numero_pdv == numero).first()
    if pdv is None:
        raise HTTPException(status_code=404, detail="Point de vente introuvable")

    type_pdv = (pdv.type_pdv or "").strip().upper() if pdv.type_pdv else None
    if type_pdv not in MIGRATION_TYPES_PDV:
        raise HTTPException(
            status_code=400,
            detail="Les appels migration concernent uniquement les PDV de type RS ou KIOSQUE.",
        )

    if not _pdv_autorise_pour_tc(db, current_user, numero):
        raise HTTPException(status_code=403, detail="Ce point de vente est affecté à une autre téléconseillère.")

    type_piece = (body.type_piece or "").strip().upper() or None
    if body.a_piece_identite and type_piece and type_piece not in MIGRATION_TYPES_PIECE:
        raise HTTPException(status_code=400, detail="Type de pièce d'identité invalide")
    if not body.a_piece_identite:
        type_piece = None

    statut, motif = _evaluer(bool(body.veut_migrer), bool(body.a_rccm), bool(body.a_piece_identite))

    m = AppelMigration(
        numero_pdv=numero,
        nom_pdv=body.nom_pdv or (pdv.nom if pdv else None),
        type_pdv=type_pdv,
        tc_user_id=current_user.id,
        tc_nom=_full_name(current_user),
        veut_migrer=bool(body.veut_migrer),
        a_rccm=bool(body.a_rccm),
        a_piece_identite=bool(body.a_piece_identite),
        type_piece=type_piece,
        statut=statut,
        motif_rejet=motif,
        commentaire=(body.commentaire or "").strip() or None,
    )
    db.add(m)
    db.commit()
    db.refresh(m)
    return _fmt(m)


class DepotBureauIn(BaseModel):
    pieces_au_bureau: bool = True


@router.patch("/tc/migration/{appel_id}/pieces-bureau")
def maj_pieces_bureau(
    appel_id: int,
    body: DepotBureauIn,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Coche (ou décoche) le dépôt des pièces au bureau pour un appel migration.

    La TC qui a passé l'appel peut le faire, ainsi que l'encadrement.
    La date du dépôt est enregistrée automatiquement lors du cochage.
    """
    m = db.query(AppelMigration).filter(AppelMigration.id == appel_id).first()
    if not m:
        raise HTTPException(status_code=404, detail="Appel migration introuvable")

    role = _role(current_user)
    if not (_est_encadrement(current_user) or m.tc_user_id == current_user.id):
        raise HTTPException(status_code=403, detail="Vous ne pouvez modifier que vos propres appels")

    m.pieces_au_bureau = bool(body.pieces_au_bureau)
    if m.pieces_au_bureau:
        m.date_depot_bureau = datetime.utcnow()
        m.depot_par_id = current_user.id
        m.depot_par_nom = _full_name(current_user)
    else:
        m.date_depot_bureau = None
        m.depot_par_id = None
        m.depot_par_nom = None

    db.commit()
    db.refresh(m)
    return _fmt(m)


# ─── Liste / statistiques / export (encadrement) ─────────────────────────────

@router.get("/tc/migration")
def list_appels_migration(
    statut: Optional[str] = Query(None, description="VALIDE ou REJETE"),
    tc_user_id: Optional[int] = Query(None),
    type_pdv: Optional[str] = Query(None, description="RS ou KIOSQUE"),
    type_piece: Optional[str] = Query(None),
    pieces_au_bureau: Optional[bool] = Query(None, description="True = pièces déposées au bureau"),
    search: Optional[str] = Query(None),
    skip: int = Query(0, ge=0),
    limit: int = Query(200, ge=1, le=1000),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Liste des appels migration (encadrement) ou de la TC connectée."""
    q = db.query(AppelMigration)
    role = _role(current_user)

    if role in ('teleconseillere', 'tc'):
        q = q.filter(AppelMigration.tc_user_id == current_user.id)
    elif not _est_encadrement(current_user):
        raise HTTPException(status_code=403, detail="Accès réservé à l'encadrement")

    if tc_user_id and role not in ('teleconseillere', 'tc'):
        q = q.filter(AppelMigration.tc_user_id == tc_user_id)
    if statut:
        q = q.filter(AppelMigration.statut == statut.strip().upper())
    if type_pdv:
        q = q.filter(AppelMigration.type_pdv == type_pdv.strip().upper())
    if type_piece:
        q = q.filter(AppelMigration.type_piece == type_piece.strip().upper())
    if pieces_au_bureau is not None:
        q = q.filter(AppelMigration.pieces_au_bureau == pieces_au_bureau)
    if search:
        like = f"%{search}%"
        q = q.filter(or_(
            AppelMigration.numero_pdv.ilike(like),
            AppelMigration.nom_pdv.ilike(like),
            AppelMigration.tc_nom.ilike(like),
        ))

    total = q.count()
    items = q.order_by(AppelMigration.created_at.desc()).offset(skip).limit(limit).all()

    # Téléphones des PDV concernés (flotte + personnel)
    tels = {}
    numeros = [m.numero_pdv for m in items]
    if numeros:
        for p in db.query(PDV).filter(PDV.numero_pdv.in_(numeros)).all():
            tels[p.numero_pdv] = (p.telephone, p.numero_personnel)

    resultat = []
    for m in items:
        d = _fmt(m)
        tel, perso = tels.get(m.numero_pdv, (None, None))
        d["telephone"] = tel
        d["numero_personnel"] = perso
        resultat.append(d)

    return {"total": total, "items": resultat}


@router.get("/tc/migration/stats")
def stats_migration(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """KPI des appels migration."""
    if not (_est_encadrement(current_user) or _role(current_user) in ('teleconseillere', 'tc')):
        raise HTTPException(status_code=403, detail="Accès refusé")

    base = db.query(AppelMigration)
    if _role(current_user) in ('teleconseillere', 'tc'):
        base = base.filter(AppelMigration.tc_user_id == current_user.id)

    total = base.count()
    valides = base.filter(AppelMigration.statut == 'VALIDE').count()
    rejetes = base.filter(AppelMigration.statut == 'REJETE').count()
    # PDV ayant accepté de migrer et PDV ayant réellement apporté leurs pièces
    acceptent = base.filter(AppelMigration.veut_migrer == True).count()
    pieces_bureau = base.filter(AppelMigration.pieces_au_bureau == True).count()

    par_tc = []
    if _est_encadrement(current_user):
        agg = {}
        for uid, nom, st, n in db.query(
            AppelMigration.tc_user_id,
            func.max(AppelMigration.tc_nom),
            AppelMigration.statut,
            func.count(AppelMigration.id),
        ).group_by(AppelMigration.tc_user_id, AppelMigration.statut).all():
            d = agg.setdefault(uid, {"tc_user_id": uid, "tc_nom": nom, "total": 0, "valides": 0,
                                     "rejetes": 0, "acceptent": 0, "pieces_bureau": 0})
            d["total"] += int(n)
            if st == 'VALIDE':
                d["valides"] += int(n)
            else:
                d["rejetes"] += int(n)
        # Détail migrer / dépôt par TC
        for uid, n in db.query(
            AppelMigration.tc_user_id, func.count(AppelMigration.id)
        ).filter(AppelMigration.veut_migrer == True).group_by(AppelMigration.tc_user_id).all():
            if uid in agg:
                agg[uid]["acceptent"] = int(n)
        for uid, n in db.query(
            AppelMigration.tc_user_id, func.count(AppelMigration.id)
        ).filter(AppelMigration.pieces_au_bureau == True).group_by(AppelMigration.tc_user_id).all():
            if uid in agg:
                agg[uid]["pieces_bureau"] = int(n)
        par_tc = sorted(agg.values(), key=lambda x: -x["total"])

    par_type = [
        {"type_pdv": t, "total": int(n)}
        for t, n in db.query(AppelMigration.type_pdv, func.count(AppelMigration.id))
        .group_by(AppelMigration.type_pdv).all()
    ]

    par_piece = [
        {"type_piece": p, "total": int(n)}
        for p, n in db.query(AppelMigration.type_piece, func.count(AppelMigration.id))
        .filter(AppelMigration.type_piece.isnot(None))
        .group_by(AppelMigration.type_piece).all()
    ]

    return {
        "total": total,
        "valides": valides,
        "rejetes": rejetes,
        "acceptent": acceptent,
        "pieces_bureau": pieces_bureau,
        "taux_depot": round(pieces_bureau / acceptent * 100, 1) if acceptent else 0,
        "taux_validation": round(valides / total * 100, 1) if total else 0,
        "par_tc": sorted(par_tc, key=lambda x: -x["total"]),
        "par_type_pdv": par_type,
        "par_piece": par_piece,
    }


@router.get("/tc/migration/export")
def export_migration(
    statut: Optional[str] = Query(None),
    tc_user_id: Optional[int] = Query(None),
    type_pdv: Optional[str] = Query(None),
    pieces_au_bureau: Optional[bool] = Query(None, description="True = uniquement les pièces déposées au bureau"),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Export Excel des appels migration (encadrement)."""
    if not _est_encadrement(current_user):
        raise HTTPException(status_code=403, detail="Export réservé à l'encadrement")

    q = db.query(AppelMigration)
    if tc_user_id:
        q = q.filter(AppelMigration.tc_user_id == tc_user_id)
    if statut:
        q = q.filter(AppelMigration.statut == statut.strip().upper())
    if type_pdv:
        q = q.filter(AppelMigration.type_pdv == type_pdv.strip().upper())
    if pieces_au_bureau is not None:
        q = q.filter(AppelMigration.pieces_au_bureau == pieces_au_bureau)
    rows = q.order_by(AppelMigration.created_at.desc()).all()

    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Appels Migration"
    entetes = [
        "#", "Date", "Téléconseillère", "N° PDV", "Nom PDV", "Type PDV",
        "Téléphone flotte", "Téléphone personnel",
        "Souhaite migrer", "RCCM", "Pièce d'identité", "Type de pièce",
        "Résultat", "Motif du rejet", "Commentaire",
        "Pièces au bureau", "Date dépôt bureau", "Dépôt enregistré par",
    ]
    ws.append(entetes)
    for c in ws[1]:
        c.font = Font(bold=True)

    # Téléphones des PDV concernés
    tels = {}
    numeros = [m.numero_pdv for m in rows]
    if numeros:
        for p in db.query(PDV).filter(PDV.numero_pdv.in_(numeros)).all():
            tels[p.numero_pdv] = (p.telephone, p.numero_personnel)

    for m in rows:
        tel, perso = tels.get(m.numero_pdv, (None, None))
        ws.append([
            m.id,
            m.created_at.strftime("%Y-%m-%d %H:%M") if m.created_at else "",
            m.tc_nom or "", m.numero_pdv, m.nom_pdv or "", m.type_pdv or "",
            tel or "", perso or "",
            "OUI" if m.veut_migrer else "NON",
            "OUI" if m.a_rccm else "NON",
            "OUI" if m.a_piece_identite else "NON",
            MIGRATION_TYPE_PIECE_LABELS.get(m.type_piece or "", "") if m.type_piece else "",
            m.statut, m.motif_rejet or "", m.commentaire or "",
            "OUI" if getattr(m, "pieces_au_bureau", False) else "NON",
            m.date_depot_bureau.strftime("%Y-%m-%d %H:%M") if getattr(m, "date_depot_bureau", None) else "",
            getattr(m, "depot_par_nom", None) or "",
        ])

    for i, largeur in enumerate([6, 17, 22, 14, 26, 10, 17, 18, 15, 8, 16, 18, 10, 45, 40, 16, 18, 22], start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = largeur

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    nom = f"appels_migration_{datetime.utcnow().strftime('%Y%m%d_%H%M')}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={nom}"},
    )
