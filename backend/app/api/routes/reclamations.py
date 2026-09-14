from fastapi import APIRouter, Depends, HTTPException, Query, UploadFile, File, Form
from fastapi.responses import FileResponse, StreamingResponse
from sqlalchemy.orm import Session
from typing import Optional, List
from datetime import datetime, timedelta
import io
import os
import uuid
from app.core.database import get_db
from app.api.routes.auth import get_current_user
from app.models.user import User
from app.models.pdv import PDV
from app.models.reclamation import (
    Reclamation, ReclamationCommentaire, ReclamationNotification,
    ReclamationHistorique, ReclamationRoutage, ReclamationPieceJointe,
)

# ─── Stockage des pièces jointes ─────────────────────────────────────────────
# Sous-dossier VOLONTAIREMENT exclu du montage statique public `/uploads`.
UPLOAD_ROOT = "uploads"
RECLAMATION_UPLOAD_SUBDIR = "reclamations"
ALLOWED_MIME = {
    "image/jpeg", "image/jpg", "image/png", "image/webp",
    "image/heic", "image/heif", "application/pdf",
}
MAX_PIECE_SIZE = 10 * 1024 * 1024  # 10 Mo

router = APIRouter()

# Routage automatique par défaut : catégorie → rôle responsable
ROUTAGE_DEFAUT = {
    'PDV': 'rc',
    'PERSONNEL': 'admin',
    'LOGISTIQUE': 'responsable_produit_et_qualit_oprationnelle_',
    'FINANCE': 'rc',
    'TECHNIQUE': 'responsable_produit_et_qualit_oprationnelle_',
    'AUTRE': 'admin',
}

# ─── Helpers rôles & autorisations ───────────────────────────────────────────

def _role(user: User) -> str:
    """Rôle normalisé (minuscules, sans préfixe d'énum)."""
    return str(user.role or '').lower().replace('userrole.', '').strip()


def _est_admin(user: User) -> bool:
    return _role(user) in ('admin', 'manager')


def _admins(db: Session):
    """Comptes admin/manager — comparaison insensible à la casse."""
    return [u for u in db.query(User).all() if _est_admin(u)]


def _peut_voir(r: Reclamation, user: User) -> bool:
    """Autorisé : admin/manager, soumetteur, ou responsable assigné."""
    return _est_admin(user) or r.soumetteur_id == user.id or r.responsable_id == user.id


def _peut_traiter(r: Reclamation, user: User) -> bool:
    """Autorisé à changer le statut / répondre : admin/manager ou responsable assigné."""
    return _est_admin(user) or r.responsable_id == user.id


def _est_en_retard(r: Reclamation) -> bool:
    """Retard calculé sur l'échéance si elle existe, sinon 72 h après création."""
    if r.statut in ('RESOLUE', 'CLOTUREE'):
        return False
    if not r.created_at:
        return False
    if r.date_limite:
        return datetime.utcnow() > r.date_limite
    return (datetime.utcnow() - r.created_at) > timedelta(hours=72)


def log_historique(db: Session, reclamation_id: int, user, action: str,
                   ancienne_valeur=None, nouvelle_valeur=None, details=None):
    """Ajoute une entrée au fil d'activité d'une réclamation."""
    nom = "Système"
    if user is not None:
        nom = f"{user.prenom or ''} {user.nom or ''}".strip() or (user.email or "Utilisateur")
    db.add(ReclamationHistorique(
        reclamation_id=reclamation_id,
        auteur_id=user.id if user is not None else None,
        auteur_nom=nom,
        action=action,
        ancienne_valeur=str(ancienne_valeur)[:200] if ancienne_valeur is not None else None,
        nouvelle_valeur=str(nouvelle_valeur)[:200] if nouvelle_valeur is not None else None,
        details=details,
    ))


# ─── Helpers ─────────────────────────────────────────────────────────────────

def notifier(db: Session, reclamation_id: int, destinataire_id: int, message: str, type_notif: str):
    notif = ReclamationNotification(
        reclamation_id=reclamation_id,
        destinataire_id=destinataire_id,
        message=message,
        type_notif=type_notif,
    )
    db.add(notif)

def notifier_admin_et_responsable(db: Session, r: Reclamation, message: str, type_notif: str, exclude_id: int = None):
    """Notifier tous les admins + le responsable assigné."""
    admins = _admins(db)
    notified_ids = set()
    for admin in admins:
        if admin.id != exclude_id:
            notifier(db, r.id, admin.id, message, type_notif)
            notified_ids.add(admin.id)
    if r.responsable_id and r.responsable_id not in notified_ids and r.responsable_id != exclude_id:
        notifier(db, r.id, r.responsable_id, message, type_notif)

def _piece_to_dict(p: ReclamationPieceJointe) -> dict:
    return {
        "id": p.id,
        "reclamation_id": p.reclamation_id,
        "auteur_nom": p.auteur_nom,
        "file_name": p.file_name,
        "mime_type": p.mime_type,
        "size_bytes": p.size_bytes or 0,
        "kind": p.kind,
        "url": f"/reclamations/{p.reclamation_id}/pieces-jointes/{p.id}/fichier",
        "created_at": p.created_at.isoformat() if p.created_at else None,
    }


def reclamation_to_dict(r: Reclamation) -> dict:
    return {
        "id": r.id,
        "titre": r.titre,
        "description": r.description,
        "categorie": r.categorie,
        "priorite": r.priorite,
        "statut": r.statut,
        "soumetteur_id": r.soumetteur_id,
        "soumetteur_nom": r.soumetteur_nom,
        "responsable_id": r.responsable_id,
        "responsable_nom": r.responsable_nom,
        "numero_pdv": r.numero_pdv,
        "nom_pdv": r.nom_pdv,
        "reponse": r.reponse,
        "date_limite": r.date_limite.isoformat() if r.date_limite else None,
        "date_prise_en_charge": r.date_prise_en_charge.isoformat() if r.date_prise_en_charge else None,
        "date_resolution": r.date_resolution.isoformat() if r.date_resolution else None,
        "note_satisfaction": r.note_satisfaction,
        "escaladee": r.escaladee,
        "escalade_raison": r.escalade_raison,
        "nb_relances": r.nb_relances,
        "created_at": r.created_at.isoformat() if r.created_at else None,
        "updated_at": r.updated_at.isoformat() if r.updated_at else None,
        # Calcul SLA
        "jours_depuis_creation": (datetime.utcnow() - r.created_at).days if r.created_at else 0,
        "en_retard": _est_en_retard(r),
    }

# ─── ROUTES RÉCLAMATIONS ─────────────────────────────────────────────────────

@router.get("/reclamations")
def list_reclamations(
    statut: Optional[str] = None,
    categorie: Optional[str] = None,
    priorite: Optional[str] = None,
    responsable_id: Optional[int] = None,
    mes_reclamations: bool = False,
    a_traiter: bool = False,
    skip: int = 0,
    limit: int = 100,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Liste des réclamations selon le rôle."""
    q = db.query(Reclamation)

    if mes_reclamations:
        q = q.filter(Reclamation.soumetteur_id == current_user.id)
    elif a_traiter:
        q = q.filter(Reclamation.responsable_id == current_user.id)
    elif not _est_admin(current_user):
        # Non-admin : voit ses réclamations + celles qu'il doit traiter
        q = q.filter(
            (Reclamation.soumetteur_id == current_user.id) |
            (Reclamation.responsable_id == current_user.id)
        )

    if statut:
        q = q.filter(Reclamation.statut == statut)
    if categorie:
        q = q.filter(Reclamation.categorie == categorie)
    if priorite:
        q = q.filter(Reclamation.priorite == priorite)
    if responsable_id:
        q = q.filter(Reclamation.responsable_id == responsable_id)

    total = q.count()
    reclamations = q.order_by(Reclamation.created_at.desc()).offset(skip).limit(limit).all()

    return {
        "total": total,
        "items": [reclamation_to_dict(r) for r in reclamations]
    }


@router.post("/reclamations")
def create_reclamation(
    data: dict,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Soumettre une nouvelle réclamation."""
    nom_complet = f"{current_user.prenom or ''} {current_user.nom or ''}".strip() or current_user.email

    titre = (data.get("titre") or "").strip()
    description = (data.get("description") or "").strip()
    if not titre or not description:
        raise HTTPException(status_code=400, detail="Titre et description sont obligatoires")

    # Mapping responsables fixes → noms affichables
    RESPONSABLES_MAP = {
        'admin': 'Admin',
        'resp_commercial': 'Responsable Commercial',
        'resp_produit': 'Responsable Produit et Qualité Opérationnelle',
        'resp_conformite': 'Resp. Conformité',
    }

    # Trouver le responsable
    responsable = None
    responsable_nom = None
    responsable_id = None
    resp_id_raw = data.get("responsable_id")
    if resp_id_raw:
        if str(resp_id_raw) in RESPONSABLES_MAP:
            # Responsable « fixe » : libellé seul, aucun compte utilisateur associé
            responsable_nom = RESPONSABLES_MAP[str(resp_id_raw)]
        else:
            try:
                responsable = db.query(User).filter(User.id == int(resp_id_raw)).first()
            except (ValueError, TypeError):
                responsable = None
            if responsable:
                responsable_id = responsable.id
                responsable_nom = f"{responsable.prenom or ''} {responsable.nom or ''}".strip()

    categorie = data.get("categorie", "AUTRE")

    # Routage automatique par catégorie si aucun responsable n'a été choisi
    routage_auto = False
    if not responsable_id and not responsable_nom:
        regle = db.query(ReclamationRoutage).filter(
            ReclamationRoutage.categorie == categorie,
            ReclamationRoutage.actif == True,
        ).first()
        if regle and regle.role_cible:
            role_cible = str(regle.role_cible).lower().replace('userrole.', '').strip()
            cible = next((u for u in db.query(User).all()
                          if str(u.role or '').lower().replace('userrole.', '').strip() == role_cible
                          and u.is_active), None)
            if cible:
                responsable = cible
                responsable_id = cible.id
                responsable_nom = f"{cible.prenom or ''} {cible.nom or ''}".strip()
                routage_auto = True

    try:
        date_limite = datetime.fromisoformat(data["date_limite"]) if data.get("date_limite") else None
    except (ValueError, TypeError):
        raise HTTPException(status_code=400, detail="Date limite invalide")

    r = Reclamation(
        titre=titre,
        description=description,
        categorie=categorie,
        priorite=data.get("priorite", "NORMAL"),
        statut="OUVERTE",
        soumetteur_id=current_user.id,
        soumetteur_nom=nom_complet,
        responsable_id=responsable_id,
        responsable_nom=responsable_nom,
        numero_pdv=data.get("numero_pdv"),
        nom_pdv=data.get("nom_pdv"),
        date_limite=date_limite,
    )
    db.add(r)
    db.flush()

    log_historique(
        db, r.id, current_user, "CREATION",
        nouvelle_valeur="OUVERTE",
        details=(f"Réclamation créée ({r.categorie} / {r.priorite})"
                 + (f" → {responsable_nom} (routage automatique)" if routage_auto
                    else (f" → {responsable_nom}" if responsable_nom else " — non assignée"))),
    )

    # Notifications
    msg_admin = f"📣 Nouvelle réclamation de {nom_complet} → {responsable_nom or 'Non assigné'} : \"{r.titre}\" [{r.priorite}]"
    msg_resp = f"📣 {nom_complet} vous a assigné une réclamation : \"{r.titre}\" [{r.priorite}]. Merci de traiter dans les 72h."

    # Notifier admins
    admins = _admins(db)
    for admin in admins:
        if admin.id != current_user.id:
            notifier(db, r.id, admin.id, msg_admin, "NOUVELLE")

    # Notifier responsable
    if responsable and responsable.id not in [a.id for a in admins]:
        notifier(db, r.id, responsable.id, msg_resp, "NOUVELLE")

    db.commit()
    db.refresh(r)
    return reclamation_to_dict(r)


@router.get("/reclamations/{rec_id}")
def get_reclamation(
    rec_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    r = db.query(Reclamation).filter(Reclamation.id == rec_id).first()
    if not r:
        raise HTTPException(status_code=404, detail="Réclamation non trouvée")
    if not _peut_voir(r, current_user):
        raise HTTPException(status_code=403, detail="Accès refusé à cette réclamation")

    # Récupérer les commentaires
    commentaires = db.query(ReclamationCommentaire).filter(
        ReclamationCommentaire.reclamation_id == rec_id
    ).order_by(ReclamationCommentaire.created_at.asc()).all()

    # Les notes internes ne sont visibles que par les admins et le responsable assigné
    peut_voir_interne = _est_admin(current_user) or r.responsable_id == current_user.id

    data = reclamation_to_dict(r)

    # Identifiant interne du PDV (pour le lien vers sa fiche)
    if r.numero_pdv:
        pdv = db.query(PDV).filter(PDV.numero_pdv == r.numero_pdv).first()
        data["pdv_id"] = pdv.id if pdv else None
    else:
        data["pdv_id"] = None

    data["commentaires"] = [{
        "id": c.id,
        "auteur_nom": c.auteur_nom,
        "auteur_role": c.auteur_role,
        "contenu": c.contenu,
        "est_interne": c.est_interne,
        "created_at": c.created_at.isoformat() if c.created_at else None,
    } for c in commentaires if peut_voir_interne or not c.est_interne]

    # Fil d'activité (traçabilité)
    historique = db.query(ReclamationHistorique).filter(
        ReclamationHistorique.reclamation_id == rec_id
    ).order_by(ReclamationHistorique.created_at.asc()).all()
    data["historique"] = [{
        "id": h.id,
        "action": h.action,
        "auteur_nom": h.auteur_nom,
        "ancienne_valeur": h.ancienne_valeur,
        "nouvelle_valeur": h.nouvelle_valeur,
        "details": h.details,
        "created_at": h.created_at.isoformat() if h.created_at else None,
    } for h in historique]

    # Pièces jointes
    pieces = db.query(ReclamationPieceJointe).filter(
        ReclamationPieceJointe.reclamation_id == rec_id
    ).order_by(ReclamationPieceJointe.id.desc()).all()
    data["pieces_jointes"] = [_piece_to_dict(p) for p in pieces]

    return data


@router.patch("/reclamations/{rec_id}")
def update_reclamation(
    rec_id: int,
    data: dict,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Mettre à jour statut, réponse, assignation, satisfaction."""
    r = db.query(Reclamation).filter(Reclamation.id == rec_id).first()
    if not r:
        raise HTTPException(status_code=404, detail="Réclamation non trouvée")
    if not _peut_voir(r, current_user):
        raise HTTPException(status_code=403, detail="Accès refusé à cette réclamation")

    # ── Contrôle des droits champ par champ ──
    est_admin = _est_admin(current_user)
    est_responsable = r.responsable_id == current_user.id
    est_soumetteur = r.soumetteur_id == current_user.id

    if "responsable_id" in data and not est_admin:
        raise HTTPException(status_code=403, detail="Seul un administrateur peut réassigner une réclamation")
    if "statut" in data:
        nouveau = data.get("statut")
        # Le soumetteur ne peut que réouvrir ; le reste est réservé au responsable/admin
        if not (est_admin or est_responsable or (est_soumetteur and nouveau == 'REOUVERTE')):
            raise HTTPException(status_code=403, detail="Seul le responsable assigné (ou un administrateur) peut changer le statut")
    if "reponse" in data and not (est_admin or est_responsable):
        raise HTTPException(status_code=403, detail="Seul le responsable assigné (ou un administrateur) peut répondre")
    if "escalade_raison" in data and not (est_admin or est_responsable):
        raise HTTPException(status_code=403, detail="Accès refusé")
    if "note_satisfaction" in data and not est_soumetteur:
        raise HTTPException(status_code=403, detail="Seul le soumetteur peut évaluer la résolution")

    nom_complet = f"{current_user.prenom or ''} {current_user.nom or ''}".strip()
    ancien_statut = r.statut
    type_notif = "MISE_A_JOUR"
    msg = None

    # Changement de statut
    if "statut" in data:
        r.statut = data["statut"]
        r.updated_at = datetime.utcnow()
        log_historique(db, r.id, current_user, "STATUT", ancienne_valeur=ancien_statut, nouvelle_valeur=data["statut"])

        if data["statut"] == "EN_COURS" and not r.date_prise_en_charge:
            r.date_prise_en_charge = datetime.utcnow()
            msg = f"✅ {nom_complet} a pris en charge votre réclamation : \"{r.titre}\""
            type_notif = "PRISE_EN_CHARGE"

        elif data["statut"] == "RESOLUE":
            r.date_resolution = datetime.utcnow()
            r.reponse = data.get("reponse", r.reponse)
            msg = f"✅ Votre réclamation \"{r.titre}\" a été résolue par {nom_complet}. Réponse : {r.reponse or '—'}"
            type_notif = "RESOLUTION"

        elif data["statut"] == "CLOTUREE":
            msg = f"🔒 Réclamation \"{r.titre}\" clôturée."
            type_notif = "CLOTURE"

        elif data["statut"] == "REOUVERTE":
            msg = f"🔄 {nom_complet} a réouvert la réclamation : \"{r.titre}\""
            type_notif = "REOUVERTURE"

        elif data["statut"] == "ESCALADEE":
            r.escaladee = True
            r.escalade_raison = data.get("escalade_raison")
            msg = f"⚠️ Réclamation escaladée par {nom_complet} : \"{r.titre}\". Raison : {r.escalade_raison}"
            type_notif = "ESCALADE"

    # Réponse sans changement de statut
    if "reponse" in data:
        r.reponse = data["reponse"]
        if data.get("statut") != "RESOLUE":
            log_historique(db, r.id, current_user, "REPONSE",
                           details=(str(data["reponse"])[:300] if data.get("reponse") else None))

    # Note de satisfaction (soumetteur uniquement)
    if "note_satisfaction" in data and current_user.id == r.soumetteur_id:
        r.note_satisfaction = data["note_satisfaction"]
        log_historique(db, r.id, current_user, "NOTE", nouvelle_valeur=data["note_satisfaction"])

    # Réassignation
    if "responsable_id" in data:
        ancien_responsable = r.responsable_nom
        if data["responsable_id"] in (None, ""):
            r.responsable_id = None
            r.responsable_nom = None
            msg = f"👤 Réclamation \"{r.titre}\" désassignée par {nom_complet}"
            type_notif = "REASSIGNATION"
        else:
            new_resp = db.query(User).filter(User.id == data["responsable_id"]).first()
            if new_resp:
                r.responsable_id = new_resp.id
                r.responsable_nom = f"{new_resp.prenom or ''} {new_resp.nom or ''}".strip()
                msg = f"👤 Réclamation \"{r.titre}\" réassignée à {r.responsable_nom} par {nom_complet}"
                type_notif = "REASSIGNATION"
        log_historique(db, r.id, current_user, "REASSIGNATION",
                       ancienne_valeur=ancien_responsable or "Non assigné",
                       nouvelle_valeur=r.responsable_nom or "Non assigné")

    # Envoyer notifications
    if msg:
        # Notifier soumetteur
        if r.soumetteur_id != current_user.id:
            notifier(db, r.id, r.soumetteur_id, msg, type_notif)
        # Notifier admins + responsable
        notifier_admin_et_responsable(db, r, msg, type_notif, exclude_id=current_user.id)

    db.commit()
    return reclamation_to_dict(r)


@router.post("/reclamations/{rec_id}/commentaires")
def add_commentaire(
    rec_id: int,
    data: dict,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    r = db.query(Reclamation).filter(Reclamation.id == rec_id).first()
    if not r:
        raise HTTPException(status_code=404, detail="Réclamation non trouvée")
    if not _peut_voir(r, current_user):
        raise HTTPException(status_code=403, detail="Accès refusé à cette réclamation")

    contenu = (data.get("contenu") or "").strip()
    if not contenu:
        raise HTTPException(status_code=400, detail="Le commentaire est vide")

    # Une note interne est réservée aux admins / responsable assigné
    est_interne = bool(data.get("est_interne", False))
    if est_interne and not (_est_admin(current_user) or r.responsable_id == current_user.id):
        est_interne = False

    nom_complet = f"{current_user.prenom or ''} {current_user.nom or ''}".strip()
    c = ReclamationCommentaire(
        reclamation_id=rec_id,
        auteur_id=current_user.id,
        auteur_nom=nom_complet,
        auteur_role=current_user.role,
        contenu=contenu,
        est_interne=est_interne,
    )
    db.add(c)

    log_historique(db, rec_id, current_user, "COMMENTAIRE",
                   details=("Note interne" if est_interne else contenu[:200]))

    # Notifier les parties prenantes
    msg = f"💬 {nom_complet} a commenté la réclamation \"{r.titre}\" : \"{contenu[:80]}...\""
    parties = set()
    if r.soumetteur_id != current_user.id:
        parties.add(r.soumetteur_id)
    if r.responsable_id and r.responsable_id != current_user.id:
        parties.add(r.responsable_id)

    for uid in parties:
        notifier(db, rec_id, uid, msg, "COMMENTAIRE")

    # Admins
    admins = _admins(db)
    for admin in admins:
        if admin.id != current_user.id and admin.id not in parties:
            notifier(db, rec_id, admin.id, msg, "COMMENTAIRE")

    db.commit()
    return {"success": True, "commentaire": {"auteur_nom": nom_complet, "contenu": contenu, "est_interne": est_interne}}


@router.post("/reclamations/{rec_id}/relancer")
def relancer_reclamation(
    rec_id: int,
    data: dict = {},
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Relancer le traitement d'une réclamation (soumetteur ou administrateur)."""
    r = db.query(Reclamation).filter(Reclamation.id == rec_id).first()
    if not r:
        raise HTTPException(status_code=404, detail="Réclamation non trouvée")
    if not _peut_voir(r, current_user):
        raise HTTPException(status_code=403, detail="Accès refusé à cette réclamation")
    if not (_est_admin(current_user) or r.soumetteur_id == current_user.id):
        raise HTTPException(status_code=403, detail="Seuls le soumetteur ou un administrateur peuvent relancer")
    if r.statut in ('RESOLUE', 'CLOTUREE'):
        raise HTTPException(status_code=400, detail="Cette réclamation est déjà traitée")

    motif = (data.get("motif") or "").strip()
    nom_complet = f"{current_user.prenom or ''} {current_user.nom or ''}".strip()
    r.nb_relances = (r.nb_relances or 0) + 1
    r.updated_at = datetime.utcnow()

    msg = f"🔔 Relance #{r.nb_relances} sur la réclamation \"{r.titre}\" par {nom_complet}"
    if motif:
        msg += f" — {motif}"

    if r.responsable_id and r.responsable_id != current_user.id:
        notifier(db, r.id, r.responsable_id, msg, "RELANCE")
    for admin in _admins(db):
        if admin.id != current_user.id and admin.id != r.responsable_id:
            notifier(db, r.id, admin.id, msg, "RELANCE")

    log_historique(db, r.id, current_user, "RELANCE",
                   nouvelle_valeur=str(r.nb_relances), details=motif or None)

    db.commit()
    return {"success": True, "nb_relances": r.nb_relances}


@router.get("/reclamations-notifications")
def get_notifications(
    non_lues_seulement: bool = False,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    q = db.query(ReclamationNotification).filter(
        ReclamationNotification.destinataire_id == current_user.id
    )
    if non_lues_seulement:
        q = q.filter(ReclamationNotification.lue == False)
    notifs = q.order_by(ReclamationNotification.created_at.desc()).limit(50).all()
    total_non_lues = db.query(ReclamationNotification).filter(
        ReclamationNotification.destinataire_id == current_user.id,
        ReclamationNotification.lue == False
    ).count()

    return {
        "total_non_lues": total_non_lues,
        "notifications": [{
            "id": n.id,
            "reclamation_id": n.reclamation_id,
            "message": n.message,
            "type_notif": n.type_notif,
            "lue": n.lue,
            "created_at": n.created_at.isoformat() if n.created_at else None,
        } for n in notifs]
    }


@router.post("/reclamations-notifications/marquer-lues")
def marquer_lues(
    data: dict = {},
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    q = db.query(ReclamationNotification).filter(
        ReclamationNotification.destinataire_id == current_user.id,
        ReclamationNotification.lue == False,
    )
    if data.get("ids"):
        q = q.filter(ReclamationNotification.id.in_(data["ids"]))
    q.update({"lue": True}, synchronize_session=False)
    db.commit()
    return {"success": True}


@router.get("/reclamations-routage")
def get_routage(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Règles de routage automatique des réclamations (catégorie → rôle responsable)."""
    rows = db.query(ReclamationRoutage).order_by(ReclamationRoutage.categorie).all()
    roles = sorted({str(u.role or '').lower().replace('userrole.', '').strip()
                    for u in db.query(User).all() if u.role})
    return {
        "regles": [{"id": r.id, "categorie": r.categorie, "role_cible": r.role_cible, "actif": bool(r.actif)} for r in rows],
        "roles_disponibles": roles,
    }


@router.put("/reclamations-routage")
def set_routage(
    data: dict,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Met à jour les règles de routage (administrateur / manager)."""
    if not _est_admin(current_user):
        raise HTTPException(status_code=403, detail="Accès réservé aux administrateurs et managers")

    for item in (data.get("regles") or []):
        cat = (item.get("categorie") or "").strip()
        if not cat:
            continue
        row = db.query(ReclamationRoutage).filter(ReclamationRoutage.categorie == cat).first()
        if not row:
            row = ReclamationRoutage(categorie=cat)
            db.add(row)
        row.role_cible = (item.get("role_cible") or "").strip().lower() or None
        row.actif = bool(item.get("actif", True))
    db.commit()
    return {"success": True}


# ─── Pièces jointes ──────────────────────────────────────────────────────────

@router.get("/reclamations/{rec_id}/pieces-jointes")
def list_pieces_jointes(
    rec_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    r = db.query(Reclamation).filter(Reclamation.id == rec_id).first()
    if not r:
        raise HTTPException(status_code=404, detail="Réclamation non trouvée")
    if not _peut_voir(r, current_user):
        raise HTTPException(status_code=403, detail="Accès refusé à cette réclamation")

    pieces = db.query(ReclamationPieceJointe).filter(
        ReclamationPieceJointe.reclamation_id == rec_id
    ).order_by(ReclamationPieceJointe.id.desc()).all()
    return {"items": [_piece_to_dict(p) for p in pieces]}


@router.post("/reclamations/{rec_id}/pieces-jointes")
async def upload_piece_jointe(
    rec_id: int,
    file: UploadFile = File(...),
    kind: str = Form("AUTRE"),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Ajouter une pièce jointe (image ou PDF, 10 Mo max)."""
    r = db.query(Reclamation).filter(Reclamation.id == rec_id).first()
    if not r:
        raise HTTPException(status_code=404, detail="Réclamation non trouvée")
    if not _peut_voir(r, current_user):
        raise HTTPException(status_code=403, detail="Accès refusé à cette réclamation")

    contenu = await file.read(MAX_PIECE_SIZE + 1)
    if len(contenu) > MAX_PIECE_SIZE:
        raise HTTPException(status_code=400, detail="Fichier trop volumineux (max 10 Mo)")
    mime = (file.content_type or "").lower()
    if mime not in ALLOWED_MIME:
        raise HTTPException(status_code=400, detail="Format non autorisé (image ou PDF uniquement)")

    dossier = os.path.join(UPLOAD_ROOT, RECLAMATION_UPLOAD_SUBDIR, str(rec_id))
    os.makedirs(dossier, exist_ok=True)
    ext = os.path.splitext(file.filename or "")[1].lower()[:10]
    nom_stocke = f"{uuid.uuid4().hex}{ext}"
    chemin = os.path.join(dossier, nom_stocke)
    with open(chemin, "wb") as f:
        f.write(contenu)

    p = ReclamationPieceJointe(
        reclamation_id=rec_id,
        auteur_id=current_user.id,
        auteur_nom=f"{current_user.prenom or ''} {current_user.nom or ''}".strip(),
        file_name=(file.filename or nom_stocke)[:300],
        file_path=chemin.replace("\\", "/"),
        mime_type=mime,
        size_bytes=len(contenu),
        kind=(kind or "AUTRE").upper()[:30],
    )
    db.add(p)
    log_historique(db, rec_id, current_user, "PIECE_JOINTE", details=p.file_name)
    db.commit()
    db.refresh(p)
    return _piece_to_dict(p)


@router.get("/reclamations/{rec_id}/pieces-jointes/{piece_id}/fichier")
def download_piece_jointe(
    rec_id: int,
    piece_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Télécharger / afficher une pièce jointe (accès authentifié obligatoire)."""
    r = db.query(Reclamation).filter(Reclamation.id == rec_id).first()
    if not r:
        raise HTTPException(status_code=404, detail="Réclamation non trouvée")
    if not _peut_voir(r, current_user):
        raise HTTPException(status_code=403, detail="Accès refusé à cette réclamation")

    p = db.query(ReclamationPieceJointe).filter(
        ReclamationPieceJointe.id == piece_id,
        ReclamationPieceJointe.reclamation_id == rec_id,
    ).first()
    if not p or not p.file_path or not os.path.exists(p.file_path):
        raise HTTPException(status_code=404, detail="Pièce jointe introuvable")

    return FileResponse(p.file_path, media_type=p.mime_type or "application/octet-stream", filename=p.file_name)


@router.delete("/reclamations/{rec_id}/pieces-jointes/{piece_id}")
def delete_piece_jointe(
    rec_id: int,
    piece_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Supprimer une pièce jointe (auteur ou administrateur)."""
    r = db.query(Reclamation).filter(Reclamation.id == rec_id).first()
    if not r:
        raise HTTPException(status_code=404, detail="Réclamation non trouvée")
    if not _peut_voir(r, current_user):
        raise HTTPException(status_code=403, detail="Accès refusé à cette réclamation")

    p = db.query(ReclamationPieceJointe).filter(
        ReclamationPieceJointe.id == piece_id,
        ReclamationPieceJointe.reclamation_id == rec_id,
    ).first()
    if not p:
        raise HTTPException(status_code=404, detail="Pièce jointe introuvable")
    if not (_est_admin(current_user) or p.auteur_id == current_user.id):
        raise HTTPException(status_code=403, detail="Seul l'auteur ou un administrateur peut supprimer cette pièce")

    nom = p.file_name
    try:
        if p.file_path and os.path.exists(p.file_path):
            os.remove(p.file_path)
    except Exception:
        pass  # le fichier a pu déjà être supprimé
    db.delete(p)
    log_historique(db, rec_id, current_user, "PIECE_JOINTE_SUPPR", details=nom)
    db.commit()
    return {"success": True}


@router.get("/reclamations-export")
def export_reclamations(
    statut: Optional[str] = None,
    categorie: Optional[str] = None,
    priorite: Optional[str] = None,
    responsable_id: Optional[int] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Export Excel des réclamations (administrateur / manager)."""
    if not _est_admin(current_user):
        raise HTTPException(status_code=403, detail="Export réservé aux administrateurs et managers")

    q = db.query(Reclamation)
    if statut:
        q = q.filter(Reclamation.statut == statut)
    if categorie:
        q = q.filter(Reclamation.categorie == categorie)
    if priorite:
        q = q.filter(Reclamation.priorite == priorite)
    if responsable_id:
        q = q.filter(Reclamation.responsable_id == responsable_id)
    rows = q.order_by(Reclamation.created_at.desc()).all()

    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Réclamations"
    entetes = ["#", "Titre", "Catégorie", "Priorité", "Statut", "Soumetteur", "Responsable",
               "PDV", "Relances", "Satisfaction", "En retard", "Créée le", "Résolue le"]
    ws.append(entetes)
    for cellule in ws[1]:
        cellule.font = Font(bold=True)

    for r in rows:
        ws.append([
            r.id, r.titre, r.categorie, r.priorite, r.statut,
            r.soumetteur_nom, r.responsable_nom or "Non assigné",
            r.numero_pdv or "", r.nb_relances or 0, r.note_satisfaction or "",
            "OUI" if _est_en_retard(r) else "non",
            r.created_at.strftime("%Y-%m-%d %H:%M") if r.created_at else "",
            r.date_resolution.strftime("%Y-%m-%d %H:%M") if r.date_resolution else "",
        ])

    largeurs = [6, 40, 12, 10, 12, 22, 22, 14, 9, 12, 10, 17, 17]
    for i, largeur in enumerate(largeurs, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = largeur

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    nom = f"reclamations_{datetime.utcnow().strftime('%Y%m%d_%H%M')}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={nom}"},
    )


@router.get("/reclamations/stats/dashboard")
def stats_dashboard(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Stats globales pour le dashboard admin."""
    if not _est_admin(current_user):
        raise HTTPException(status_code=403, detail="Accès réservé aux administrateurs et managers")
    all_rec = db.query(Reclamation).all()
    maintenant = datetime.utcnow()

    # Délai moyen de résolution (heures)
    res_avec_delai = [r for r in all_rec if r.date_resolution and r.created_at]
    delai_moyen_h = round(
        sum((r.date_resolution - r.created_at).total_seconds() for r in res_avec_delai)
        / len(res_avec_delai) / 3600, 1
    ) if res_avec_delai else 0

    # Satisfaction moyenne
    notes = [r.note_satisfaction for r in all_rec if r.note_satisfaction]
    satisfaction = round(sum(notes) / len(notes), 2) if notes else 0

    # Respect du délai (SLA) sur les réclamations traitées
    traitees = [r for r in all_rec if r.statut in ('RESOLUE', 'CLOTUREE')]

    def _dans_sla(r):
        if not r.created_at:
            return True
        limite = r.date_limite or (r.created_at + timedelta(hours=72))
        fin = r.date_resolution or maintenant
        return fin <= limite

    taux_sla = round(sum(1 for r in traitees if _dans_sla(r)) / len(traitees) * 100) if traitees else 0

    # Répartition par responsable
    par_resp = {}
    for r in all_rec:
        nom = r.responsable_nom or 'Non assigné'
        d = par_resp.setdefault(nom, {'responsable_nom': nom, 'total': 0, 'resolues': 0, 'en_retard': 0})
        d['total'] += 1
        if r.statut in ('RESOLUE', 'CLOTUREE'):
            d['resolues'] += 1
        if _est_en_retard(r):
            d['en_retard'] += 1
    par_responsable = sorted(par_resp.values(), key=lambda x: -x['total'])

    return {
        "total": len(all_rec),
        "ouvertes": sum(1 for r in all_rec if r.statut == "OUVERTE"),
        "en_cours": sum(1 for r in all_rec if r.statut == "EN_COURS"),
        "resolues": sum(1 for r in all_rec if r.statut == "RESOLUE"),
        "cloturees": sum(1 for r in all_rec if r.statut == "CLOTUREE"),
        "en_retard": sum(1 for r in all_rec if _est_en_retard(r)),
        "urgentes": sum(1 for r in all_rec if r.priorite == "URGENT" and r.statut not in ('RESOLUE','CLOTUREE')),
        "escaladees": sum(1 for r in all_rec if r.escaladee),
        "relancees": sum(1 for r in all_rec if (r.nb_relances or 0) > 0),
        "non_assignees": sum(1 for r in all_rec if not r.responsable_id and r.statut not in ('RESOLUE','CLOTUREE')),
        "taux_resolution": round(sum(1 for r in all_rec if r.statut in ('RESOLUE','CLOTUREE')) / len(all_rec) * 100) if all_rec else 0,
        "taux_sla": taux_sla,
        "delai_moyen_resolution_h": delai_moyen_h,
        "satisfaction_moyenne": satisfaction,
        "par_categorie": {cat: sum(1 for r in all_rec if r.categorie == cat) for cat in ['PDV','PERSONNEL','LOGISTIQUE','FINANCE','TECHNIQUE','AUTRE']},
        "par_priorite": {p: sum(1 for r in all_rec if r.priorite == p) for p in ['URGENT','NORMAL','FAIBLE']},
        "par_responsable": par_responsable,
    }
