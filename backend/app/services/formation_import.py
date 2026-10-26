"""
Import du classeur Excel « Suivi Formation » vers la base.

Le fichier historique contient trois feuilles :

  • BASE DE DONNEES   → référentiel des agents
  • SUIVI FORMATIONS  → une colonne de statut + une colonne de date par module
  • FICHE INDIVIDUELLE → formule de consultation (non importée : elle est recalculée)

L'import est volontairement tolérant : il ne rejette pas le fichier à cause de
cellules approximatives, il corrige ce qui est réparable et **rapporte** tout ce
qu'il a dû interpréter (l'utilisateur garde la maîtrise).
"""
from __future__ import annotations

import io
import re
import unicodedata
from datetime import date, datetime
from typing import Any, Dict, List, Optional, Tuple

import openpyxl
from sqlalchemy.orm import Session

from app.models.formation import FormationModule, FormationPersonnel, FormationSession

# Les 8 modules historiques (feuilles 1 : colonnes F, H, J, L, N, P, R, T)
MODULES_DEFAUT = [
    "L'AGENCE / POINT HUB OMY",
    "OBLIGATIONS PRINCIPALES D'UN SUPERVISEUR",
    "BESOINS D'UN GESTIONNAIRE OMY & IDENTIFICATION",
    "RÔLE DU SUPERVISEUR OMY & IDENTIFICATION",
    "KAABU +",
    "ACTIVATIONS SIM PDV",
    "RÔLE D'UNE TÉLÉCONSEILLÈRE OMY & IDENTIFICATION",
    "ACTIVITÉS BG",
]

STATUTS = {"A_FORMER", "EN_COURS", "VALIDER", "A_REFORMER", "NA"}
EVALUATIONS = {"EXCELLENT", "TRES_BIEN", "BIEN", "MOYEN", "INSUFFISANT"}

_RE_DATES = re.compile(r"(\d{1,2})\s*[/\-.]\s*(\d{1,2})\s*[/\-.]\s*(\d{2,4})")
_RE_NOTE = re.compile(r"(\d{1,2}(?:[.,]\d+)?)\s*/\s*20")


# ── Utilitaires de normalisation ──────────────────────────────────────────────
def sans_accents(texte: Any) -> str:
    if texte is None:
        return ""
    t = unicodedata.normalize("NFKD", str(texte))
    t = "".join(c for c in t if not unicodedata.combining(c))
    return " ".join(t.upper().split())


def _texte(v: Any) -> Optional[str]:
    if v is None:
        return None
    t = str(v).strip()
    return t or None


def normaliser_statut(v: Any) -> Optional[str]:
    """« VALIDER EN LIGNE » → VALIDER ; « À REFORMER » → A_REFORMER ; etc."""
    s = sans_accents(v)
    if not s or s in {"NONE", "NAN"}:
        return None
    if s.startswith("VALID"):
        return "VALIDER"
    table = {
        "A FORMER": "A_FORMER",
        "EN COURS": "EN_COURS",
        "A REFORMER": "A_REFORMER",
        "NA": "NA",
        "N/A": "NA",
    }
    return table.get(s, s if s in STATUTS else None)


def normaliser_evaluation(v: Any) -> Optional[str]:
    s = sans_accents(v)
    table = {
        "EXCELLENT": "EXCELLENT",
        "TRES BIEN": "TRES_BIEN",
        "BIEN": "BIEN",
        "MOYEN": "MOYEN",
        "INSUFFISANT": "INSUFFISANT",
    }
    return table.get(s)


def extraire_dates(v: Any) -> Tuple[List[date], Optional[str]]:
    """Retourne (dates valides trouvées, texte d'origine si non exploitable).

    Une cellule peut contenir plusieurs formations : « 15/10/2024 17/01/2025 ».
    """
    if v is None:
        return [], None
    if isinstance(v, datetime):
        return [v.date()], None
    if isinstance(v, date):
        return [v], None

    texte = str(v).strip()
    if not texte:
        return [], None

    dates: List[date] = []
    for jour, mois, annee in _RE_DATES.findall(texte):
        a = int(annee)
        if a < 100:
            a += 2000
        try:
            dates.append(date(a, int(mois), int(jour)))
        except ValueError:
            continue

    toutes_numbers = bool(dates) and len(dates) == len(_RE_DATES.findall(texte))
    return dates, (None if toutes_numbers else texte)


def _telephone(v: Any) -> Optional[str]:
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, int):
        return str(v)
    t = str(v).strip()
    return t or None


# ── Repérage des feuilles et des colonnes ─────────────────────────────────────
def _feuille(wb, nom_attendu: str):
    cible = sans_accents(nom_attendu)
    for ws in wb.worksheets:
        if sans_accents(ws.title) == cible:
            return ws
    # tolérance : recherche partielle
    for ws in wb.worksheets:
        if cible in sans_accents(ws.title) or sans_accents(ws.title) in cible:
            return ws
    return None


def _ligne_entete(ws, libelles: List[str], limite: int = 15) -> Optional[int]:
    cibles = [sans_accents(l) for l in libelles]
    for r in range(1, min(limite, ws.max_row or 1) + 1):
        valeurs = {sans_accents(ws.cell(row=r, column=c).value) for c in range(1, (ws.max_column or 0) + 1)}
        if all(c in valeurs for c in cibles):
            return r
    return None


def _colonnes(ws, ligne: int) -> Dict[str, int]:
    mapping: Dict[str, int] = {}
    for c in range(1, (ws.max_column or 0) + 1):
        cle = sans_accents(ws.cell(row=ligne, column=c).value)
        if cle and cle not in mapping:
            mapping[cle] = c
    return mapping


def _col(mapping: Dict[str, int], *libelles: str) -> Optional[int]:
    for libelle in libelles:
        c = mapping.get(sans_accents(libelle))
        if c:
            return c
    return None


# ── Import ───────────────────────────────────────────────────────────────────
def importer_classeur(contenu: bytes, db: Session, user_id: Optional[int] = None) -> Dict[str, Any]:
    """Importe le classeur et renvoie un rapport détaillé."""
    try:
        wb = openpyxl.load_workbook(io.BytesIO(contenu), data_only=True)
    except Exception as exc:  # fichier illisible
        raise ValueError(f"Fichier Excel illisible : {exc}") from exc

    ws_base = _feuille(wb, "BASE DE DONNEES")
    ws_suivi = _feuille(wb, "SUIVI FORMATIONS")
    if ws_base is None:
        raise ValueError("La feuille « BASE DE DONNEES » est introuvable dans le fichier.")

    rapport: Dict[str, Any] = {
        "ok": True,
        "personnel_crees": 0,
        "personnel_mis_a_jour": 0,
        "personnel_ignores": 0,
        "modules_crees": 0,
        "sessions_creees": 0,
        "sessions_ignorees": 0,
        "evaluations_reprises": 0,
        "avertissements": [],
    }
    avert = rapport["avertissements"]

    def alerte(message: str):
        if len(avert) < 200:
            avert.append(message)

    # ── 1. Catalogue des modules (créés si absents) ──────────────────────────
    modules: Dict[str, FormationModule] = {}
    for m in db.query(FormationModule).all():
        modules[sans_accents(m.nom)] = m

    for ordre, nom in enumerate(MODULES_DEFAUT):
        cle = sans_accents(nom)
        if cle in modules:
            continue
        m = FormationModule(nom=nom, ordre=ordre, obligatoire=True, actif=True)
        db.add(m)
        db.flush()
        modules[cle] = m
        rapport["modules_crees"] += 1

    # ── 2. Personnel ─────────────────────────────────────────────────────────
    ligne = _ligne_entete(ws_base, ["ID PERSONNEL", "NOM ET PRENOM"])
    if ligne is None:
        raise ValueError("En-têtes « ID PERSONNEL » / « NOM ET PRÉNOM » introuvables dans BASE DE DONNEES.")
    colonnes = _colonnes(ws_base, ligne)

    c_id = _col(colonnes, "ID PERSONNEL")
    c_nom = _col(colonnes, "NOM ET PRENOM")
    c_tel = _col(colonnes, "N° TELEPHONE", "N TELEPHONE", "TELEPHONE")
    c_gest = _col(colonnes, "GESTIONNAIRE")
    c_groupe = _col(colonnes, "GROUPE")
    c_poste = _col(colonnes, "POSTE")
    c_statut = _col(colonnes, "STATUT")

    if not c_id or not c_nom:
        raise ValueError("Colonnes « ID PERSONNEL » ou « NOM ET PRÉNOM » absentes.")

    par_matricule: Dict[str, FormationPersonnel] = {
        (p.matricule or "").strip().upper(): p for p in db.query(FormationPersonnel).all()
    }
    postes_deplaces = 0

    for r in range(ligne + 1, (ws_base.max_row or 0) + 1):
        matricule = _texte(ws_base.cell(row=r, column=c_id).value)
        nom = _texte(ws_base.cell(row=r, column=c_nom).value)
        if not matricule and not nom:
            continue
        if not nom:
            rapport["personnel_ignores"] += 1
            alerte(f"Ligne {r} : nom manquant (matricule {matricule or 'absent'}) — ignorée.")
            continue
        if not matricule:
            matricule = f"SANS-ID-{r}"
            alerte(f"Ligne {r} : {nom} n'a pas de matricule — identifiant provisoire {matricule}.")

        telephone = _telephone(ws_base.cell(row=r, column=c_tel).value) if c_tel else None
        gestionnaire = _texte(ws_base.cell(row=r, column=c_gest).value) if c_gest else None
        groupe = _texte(ws_base.cell(row=r, column=c_groupe).value) if c_groupe else None
        poste = _texte(ws_base.cell(row=r, column=c_poste).value) if c_poste else None
        statut = _texte(ws_base.cell(row=r, column=c_statut).value) if c_statut else None

        # 31 lignes du fichier portent l'adresse du Point Hub dans la colonne POSTE
        agence = None
        if poste and sans_accents(poste).startswith("N°"):
            agence = poste
            poste = "POINT HUB"
            postes_deplaces += 1
            alerte(f"{nom} ({matricule}) : adresse trouvée dans POSTE — déplacée vers AGENCE, poste = POINT HUB.")

        existant = par_matricule.get(matricule.strip().upper())
        if existant:
            existant.nom = nom
            existant.telephone = telephone or existant.telephone
            existant.gestionnaire = gestionnaire or existant.gestionnaire
            existant.groupe = groupe or existant.groupe
            existant.poste = poste or existant.poste
            existant.agence = agence or existant.agence
            existant.statut_emploi = statut or existant.statut_emploi
            existant.actif = True
            rapport["personnel_mis_a_jour"] += 1
        else:
            p = FormationPersonnel(
                matricule=matricule.strip(),
                nom=nom,
                telephone=telephone,
                gestionnaire=gestionnaire,
                groupe=groupe,
                poste=poste,
                agence=agence,
                statut_emploi=statut or "ACTIF",
                actif=True,
            )
            db.add(p)
            db.flush()
            par_matricule[matricule.strip().upper()] = p
            rapport["personnel_crees"] += 1

    if postes_deplaces:
        alerte(f"{postes_deplaces} agent(s) avaient une adresse à la place du poste (voir ci-dessus).")

    db.flush()

    # ── 3. Formations (feuille SUIVI FORMATIONS) ─────────────────────────────
    if ws_suivi is not None:
        ligne_s = _ligne_entete(ws_suivi, ["ID PERSONNEL", "NOM ET PRENOM"])
        if ligne_s is None:
            alerte("Feuille SUIVI FORMATIONS : en-têtes introuvables — formations non importées.")
        else:
            cols_s = _colonnes(ws_suivi, ligne_s)
            cs_id = _col(cols_s, "ID PERSONNEL")
            entetes = {c: sans_accents(ws_suivi.cell(row=ligne_s, column=c).value)
                       for c in range(1, (ws_suivi.max_column or 0) + 1)}

            # Un module = une colonne dont la SUIVANTE s'appelle « DATE FORMATION »
            blocs = []
            for c in range(1, (ws_suivi.max_column or 0)):
                if entetes.get(c + 1) == "DATE FORMATION" and entetes.get(c):
                    blocs.append((entetes[c], c, c + 1))

            col_eval = next((c for c, e in entetes.items()
                             if e and ("EVALUATION" in e or "COMPREHENSION" in e)), None)

            if not blocs:
                alerte("Aucune colonne de module détectée dans SUIVI FORMATIONS.")

            for r in range(ligne_s + 1, (ws_suivi.max_row or 0) + 1):
                matricule = _texte(ws_suivi.cell(row=r, column=cs_id).value) if cs_id else None
                if not matricule:
                    continue
                personne = par_matricule.get(matricule.strip().upper())
                if personne is None:
                    rapport["personnel_ignores"] += 1
                    alerte(f"SUIVI ligne {r} : matricule {matricule} absent de la base — formations ignorées.")
                    continue

                for nom_module, c_statut_xlsx, c_date_xlsx in blocs:
                    brut_statut = ws_suivi.cell(row=r, column=c_statut_xlsx).value
                    brut_date = ws_suivi.cell(row=r, column=c_date_xlsx).value
                    statut = normaliser_statut(brut_statut)
                    dates, texte_brut = extraire_dates(brut_date)

                    if statut is None and not dates:
                        continue
                    if statut is None:
                        statut = "VALIDER"
                        alerte(f"{personne.nom} · {nom_module} : date sans statut — validé par défaut.")

                    module = modules.get(sans_accents(nom_module))
                    if module is None:
                        module = FormationModule(nom=nom_module, ordre=len(modules), obligatoire=False, actif=True)
                        db.add(module)
                        db.flush()
                        modules[sans_accents(nom_module)] = module
                        rapport["modules_crees"] += 1

                    date_ret = max(dates) if dates else None
                    observation = None
                    if texte_brut:
                        observation = f"Date(s) d'origine dans le fichier : {texte_brut}"
                        alerte(f"{personne.nom} · {nom_module} : date non exploitable « {texte_brut} ».")
                    elif len(dates) > 1:
                        observation = "Date(s) d'origine : " + " ".join(d.strftime("%d/%m/%Y") for d in dates)
                        alerte(f"{personne.nom} · {nom_module} : {len(dates)} dates dans une cellule — la plus récente retenue.")

                    existe = db.query(FormationSession).filter(
                        FormationSession.personnel_id == personne.id,
                        FormationSession.module_id == module.id,
                        FormationSession.date_session == date_ret,
                    ).first()
                    if existe:
                        rapport["sessions_ignorees"] += 1
                        continue

                    db.add(FormationSession(
                        personnel_id=personne.id,
                        module_id=module.id,
                        date_session=date_ret,
                        statut=statut,
                        observation=observation,
                        cree_par_id=user_id,
                    ))
                    rapport["sessions_creees"] += 1

                # Évaluation globale (« BIEN | 18/20 »)
                if col_eval:
                    brut = ws_suivi.cell(row=r, column=col_eval).value
                    if brut:
                        texte = str(brut)
                        evaluation = normaliser_evaluation(texte.split("|")[0])
                        note = None
                        m_note = _RE_NOTE.search(texte)
                        if m_note:
                            note = float(m_note.group(1).replace(",", "."))
                        if note is None and "|" in texte:
                            try:
                                note = float(texte.split("|")[1].strip())
                            except (ValueError, IndexError):
                                note = None
                        if evaluation or note is not None:
                            personne.evaluation_globale = evaluation or personne.evaluation_globale
                            personne.note_globale = note if note is not None else personne.note_globale
                            rapport["evaluations_reprises"] += 1

    db.commit()
    return rapport


def assurer_modules_defaut(db: Session) -> int:
    """Crée les modules historiques manquants (idempotent).

    Appelé au démarrage : le catalogue est ainsi disponible avant même le
    premier import, et l'utilisateur peut le modifier librement ensuite.
    """
    existants = {sans_accents(m.nom) for m in db.query(FormationModule).all()}
    crees = 0
    for ordre, nom in enumerate(MODULES_DEFAUT):
        if sans_accents(nom) in existants:
            continue
        db.add(FormationModule(nom=nom, ordre=ordre, obligatoire=True, actif=True))
        crees += 1
    if crees:
        db.commit()
    return crees
