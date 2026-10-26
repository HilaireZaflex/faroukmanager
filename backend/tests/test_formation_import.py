"""
Tests du module « Suivi Formation » : import du classeur historique.

Un classeur synthétique reproduit la structure réelle du fichier Farouk
Distribution (3 feuilles, en-têtes décalés, dates multiples dans une cellule,
adresse dans la colonne POSTE, évaluations « BIEN | 18/20 »).

Aucune donnée client n'est utilisée : le test est autonome.
"""
import io

import pytest
from openpyxl import Workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.core.database import Base
from app.models.formation import FormationModule, FormationPersonnel, FormationSession
from app.models.user import User
from app.services.formation_import import MODULES_DEFAUT, assurer_modules_defaut, importer_classeur


# ── Classeur de test ─────────────────────────────────────────────────────────
def _classeur() -> bytes:
    wb = Workbook()

    base = wb.active
    base.title = "BASE DE DONNEES"
    base["A1"] = "BASE DE DONNÉES – PERSONNEL"
    base.append([])
    base.append(["ID PERSONNEL", "NOM ET PRÉNOM", "N° TÉLÉPHONE", "GESTIONNAIRE",
                 "GROUPE", "POSTE", "STATUT"])
    base.append(["FRK-001", "ALPHA BETA", 70000001, "GEST X", "BG", "SUPERVISEUR", "ACTIF"])
    base.append(["FRK-002", "GAMMA DELTA", 70000002, "GEST X", "BG",
                 "N°9 POINT HUB TEST", "ACTIF"])
    base.append(["FRK-003", "EPSILON ZETA", 70000003, "GEST Y", "GRP 1", "TELECONSEILLERE", "ACTIF"])

    suivi = wb.create_sheet("SUIVI FORMATIONS")
    suivi.append(["TABLEAU DE SUIVI"])
    suivi.append([])
    suivi.append([])
    suivi.append(["ID PERSONNEL", "NOM ET PRÉNOM", "N° TÉLÉPHONE", "GROUPE", "POSTE",
                  "L'AGENCE / POINT HUB OMY", "DATE FORMATION",
                  "OBLIGATIONS PRINCIPALES D'UN SUPERVISEUR", "DATE FORMATION",
                  "COMPRÉHENSION / ÉVALUATION"])
    suivi.append(["FRK-001", "ALPHA BETA", 70000001, "BG", "SUPERVISEUR",
                  "VALIDER", "15/10/2024 17/01/2025", "VALIDER", "17/06/2025", "BIEN | 18/20"])
    suivi.append(["FRK-002", "GAMMA DELTA", 70000002, "BG", "POINT HUB", None, None, None, None, None])
    suivi.append(["FRK-003", "EPSILON ZETA", 70000003, "GRP 1", "TELECONSEILLERE",
                  "VALIDER EN LIGNE", None, "À REFORMER", "17/010/2025", "MOYEN | 12/20"])

    wb.create_sheet("FICHE INDIVIDUELLE")

    flux = io.BytesIO()
    wb.save(flux)
    return flux.getvalue()


@pytest.fixture()
def db():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine, tables=[
        User.__table__, FormationPersonnel.__table__,
        FormationModule.__table__, FormationSession.__table__,
    ])
    session = sessionmaker(bind=engine)()
    yield session
    session.close()


@pytest.fixture()
def rapport(db):
    return importer_classeur(_classeur(), db)


# ── Tests ────────────────────────────────────────────────────────────────────
def test_personnel_importe(rapport, db):
    assert rapport["personnel_crees"] == 3
    assert db.query(FormationPersonnel).count() == 3

    alpha = db.query(FormationPersonnel).filter(FormationPersonnel.matricule == "FRK-001").one()
    assert alpha.nom == "ALPHA BETA"
    assert alpha.telephone == "70000001"
    assert alpha.gestionnaire == "GEST X"
    assert alpha.evaluation_globale == "BIEN"
    assert alpha.note_globale == 18.0


def test_catalogue_modules_cree(rapport, db):
    assert rapport["modules_crees"] == len(MODULES_DEFAUT)
    assert db.query(FormationModule).count() == len(MODULES_DEFAUT)
    noms = {m.nom for m in db.query(FormationModule).all()}
    assert "KAABU +" in noms


def test_adresse_dans_colonne_poste_deplacee(rapport, db):
    gamma = db.query(FormationPersonnel).filter(FormationPersonnel.matricule == "FRK-002").one()
    assert gamma.poste == "POINT HUB"
    assert gamma.agence == "N°9 POINT HUB TEST"
    assert any("adresse trouvée dans POSTE" in a for a in rapport["avertissements"])


def test_dates_multiples_garde_la_plus_recente(rapport, db):
    alpha = db.query(FormationPersonnel).filter(FormationPersonnel.matricule == "FRK-001").one()
    module = db.query(FormationModule).filter(
        FormationModule.nom == "L'AGENCE / POINT HUB OMY").one()
    s = db.query(FormationSession).filter(
        FormationSession.personnel_id == alpha.id,
        FormationSession.module_id == module.id).one()

    assert s.statut == "VALIDER"
    assert s.date_session.isoformat() == "2025-01-17"
    assert "15/10/2024" in (s.observation or "")


def test_statuts_normalises(rapport, db):
    epsilon = db.query(FormationPersonnel).filter(FormationPersonnel.matricule == "FRK-003").one()
    agence = db.query(FormationModule).filter(
        FormationModule.nom == "L'AGENCE / POINT HUB OMY").one()
    s = db.query(FormationSession).filter(
        FormationSession.personnel_id == epsilon.id,
        FormationSession.module_id == agence.id).one()

    assert s.statut == "VALIDER"          # « VALIDER EN LIGNE »
    assert s.date_session is None          # date absente

    obligations = db.query(FormationModule).filter(
        FormationModule.nom == "OBLIGATIONS PRINCIPALES D'UN SUPERVISEUR").one()
    s2 = db.query(FormationSession).filter(
        FormationSession.personnel_id == epsilon.id,
        FormationSession.module_id == obligations.id).one()
    assert s2.statut == "A_REFORMER"       # « À REFORMER »
    assert s2.date_session is None         # « 17/010/2025 » inexploitable
    assert "17/010/2025" in (s2.observation or "")


def test_evaluation_moyenne_reprise(rapport, db):
    epsilon = db.query(FormationPersonnel).filter(FormationPersonnel.matricule == "FRK-003").one()
    assert epsilon.evaluation_globale == "MOYEN"
    assert epsilon.note_globale == 12.0


def test_reimport_ne_duplique_rien(db):
    premiere = importer_classeur(_classeur(), db)
    nb_sessions = db.query(FormationSession).count()
    assert premiere["sessions_creees"] > 0

    seconde = importer_classeur(_classeur(), db)
    assert seconde["personnel_crees"] == 0
    assert seconde["sessions_creees"] == 0
    assert db.query(FormationSession).count() == nb_sessions
    assert db.query(FormationPersonnel).count() == 3


def test_fichier_sans_feuille_base_rejetee(db):
    wb = Workbook()
    wb.active.title = "AUTRE"
    flux = io.BytesIO()
    wb.save(flux)
    with pytest.raises(ValueError):
        importer_classeur(flux.getvalue(), db)


def test_initialisation_catalogue_idempotente(db):
    assert assurer_modules_defaut(db) == len(MODULES_DEFAUT)
    assert assurer_modules_defaut(db) == 0
    assert db.query(FormationModule).count() == len(MODULES_DEFAUT)


# ── Matrice de suivi (lot 2) ─────────────────────────────────────────────────
from app.api.routes.formations import matrice, detail_personnel  # noqa: E402


class _Admin:
    role = "admin"
    id = 1


def test_matrice_retourne_modules_et_cellules(rapport, db):
    r = matrice(search=None, groupe=None, poste=None, gestionnaire=None,
                actif=True, db=db, current_user=_Admin())

    assert len(r["modules"]) == len(MODULES_DEFAUT)
    assert r["total"] == 3

    alpha = next(p for p in r["personnes"] if p["matricule"] == "FRK-001")
    module_agence = next(m for m in r["modules"] if m["nom"] == "L'AGENCE / POINT HUB OMY")
    cellule = alpha["cellules"].get(str(module_agence["id"]))

    assert cellule is not None
    assert cellule["statut"] == "VALIDER"
    assert cellule["date_session"] == "2025-01-17"
    assert alpha["nb_valides"] == 2          # agence + obligations


def test_matrice_filtre_par_groupe(rapport, db):
    r = matrice(search=None, groupe="GRP 1", poste=None, gestionnaire=None,
                actif=True, db=db, current_user=_Admin())
    assert r["total"] == 1
    assert r["personnes"][0]["matricule"] == "FRK-003"


def test_matrice_recherche(rapport, db):
    r = matrice(search="GAMMA", groupe=None, poste=None, gestionnaire=None,
                actif=True, db=db, current_user=_Admin())
    assert r["total"] == 1
    assert r["personnes"][0]["nom"] == "GAMMA DELTA"


def test_detail_personnel_avec_historique(rapport, db):
    alpha = db.query(FormationPersonnel).filter(FormationPersonnel.matricule == "FRK-001").one()
    d = detail_personnel(personnel_id=alpha.id, db=db, current_user=_Admin())

    assert d["nom"] == "ALPHA BETA"
    assert len(d["historique"]) == 2
    assert d["historique"][0]["module_nom"]
    assert d["nb_valides"] == 2


def test_matrice_refusee_pour_un_role_non_autorise(rapport, db):
    class _TC:
        role = "teleconseillere"
        id = 9

    with pytest.raises(Exception) as err:
        matrice(search=None, groupe=None, poste=None, gestionnaire=None,
                actif=True, db=db, current_user=_TC())
    assert "403" in str(err.value) or "réservé" in str(err.value)


def test_fiche_liste_tous_les_modules(rapport, db):
    alpha = db.query(FormationPersonnel).filter(FormationPersonnel.matricule == "FRK-001").one()
    d = detail_personnel(personnel_id=alpha.id, db=db, current_user=_Admin())

    assert len(d["parcours"]) == len(MODULES_DEFAUT)
    agence = next(x for x in d["parcours"] if x["module_nom"] == "L'AGENCE / POINT HUB OMY")
    assert agence["statut"] == "VALIDER"
    assert agence["date_session"] == "2025-01-17"

    sans = [x for x in d["parcours"] if x["statut"] is None]
    assert d["nb_non_formes"] == len(sans)
    assert len(sans) == len(MODULES_DEFAUT) - 2
    assert d["alerte"] == "INCOMPLET"      # 2 modules validés sur 8


def test_fiche_alerte_action_requise(rapport, db):
    epsilon = db.query(FormationPersonnel).filter(FormationPersonnel.matricule == "FRK-003").one()
    d = detail_personnel(personnel_id=epsilon.id, db=db, current_user=_Admin())
    assert d["alerte"] == "ACTION_REQUISE"
    reformer = [x for x in d["parcours"] if x["statut"] == "A_REFORMER"]
    assert len(reformer) == 1


# ── Pilotage & export (lot 4) ───────────────────────────────────────────────
import asyncio  # noqa: E402
import io as _io  # noqa: E402

import openpyxl  # noqa: E402

from app.api.routes.formations import dashboard, exporter_excel, lister_alertes  # noqa: E402


def test_dashboard_kpis_et_modules(rapport, db):
    d = dashboard(groupe=None, poste=None, db=db, current_user=_Admin())

    assert d["kpis"]["nb_personnel"] == 3
    assert d["kpis"]["nb_modules"] == len(MODULES_DEFAUT)
    assert len(d["par_module"]) == len(MODULES_DEFAUT)
    assert sum(d["alertes"].values()) == 3

    agence = next(m for m in d["par_module"] if m["nom"] == "L'AGENCE / POINT HUB OMY")
    assert agence["nb_valides"] == 2          # FRK-001 et FRK-003
    assert agence["nb_non_formes"] == 1       # FRK-002


def test_alertes_filtrees_par_type(rapport, db):
    toutes = lister_alertes(type_alerte=None, groupe=None, poste=None, gestionnaire=None,
                            db=db, current_user=_Admin())
    assert toutes["total"] == 3
    assert toutes["compteurs"]["ACTION_REQUISE"] == 1     # FRK-003

    action = lister_alertes(type_alerte="ACTION_REQUISE", groupe=None, poste=None,
                            gestionnaire=None, db=db, current_user=_Admin())
    assert action["total"] == 1
    assert action["items"][0]["matricule"] == "FRK-003"
    assert action["items"][0]["modules_a_reformer"]
    assert action["items"][0]["modules_manquants"]


def test_export_excel_quatre_feuilles(rapport, db):
    resp = exporter_excel(groupe=None, poste=None, db=db, current_user=_Admin())

    async def collecter(r):
        return b"".join([bloc async for bloc in r.body_iterator])

    contenu = asyncio.run(collecter(resp))
    wb = openpyxl.load_workbook(_io.BytesIO(contenu))

    assert wb.sheetnames == ["Synthèse", "Matrice", "Détail sessions", "Personnel"]

    matrice = wb["Matrice"]
    entetes = [matrice.cell(row=1, column=c).value for c in range(1, 8)]
    assert entetes[:6] == ["ID PERSONNEL", "NOM ET PRÉNOM", "N° TÉLÉPHONE",
                           "GROUPE", "POSTE", "AGENCE"]
    assert entetes[6] == "L'AGENCE / POINT HUB OMY"
    assert matrice.cell(row=1, column=7 + 2 * len(MODULES_DEFAUT)).value == "COMPRÉHENSION / ÉVALUATION"

    # le vocabulaire d'origine est conservé
    valeurs = [matrice.cell(row=r, column=7).value for r in range(2, 4)]
    assert all(v in (None, "VALIDER", "À FORMER", "EN COURS", "À REFORMER", "N/A") for v in valeurs)
